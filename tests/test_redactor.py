"""Unit tests for run preservation, caching, family consistency, regex, and OCR redaction."""
from __future__ import annotations

import io
import pytest
from docx import Document
from PIL import Image, ImageDraw, ImageFont

from pii_redactor import (
    Anonymizer,
    DocxEngine,
    EntityType as T,
    ImageRedactor,
    Pseudonymizer,
    RedactionConfig,
    Redactor,
)
from pii_redactor.detector import PIIDetector
from pii_redactor.patterns import (
    AADHAAR_RE,
    CIN_RE,
    DIN_RE,
    EMAIL_RE,
    PAN_RE,
    PHONE_INTL_RE,
    PHONE_LANDLINE_RE,
    PHONE_MOBILE_RE,
    SEBI_REG_RE,
)
from pii_redactor.validators import luhn_ok, verhoeff_ok


def _doc_bytes(build) -> bytes:
    d = Document()
    build(d)
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


# ---------------------------------------------------------------------------
# 1. Run Preservation Tests
# ---------------------------------------------------------------------------
def test_split_runs_preserve_individual_formatting():
    """Ensures replacement across runs preserves each run's distinct bold/italic properties."""
    def build(d):
        p = d.add_paragraph()
        p.add_run("Promoter: ")
        r1 = p.add_run("Kus")
        r1.bold = True
        r2 = p.add_run("hal Subbayya ")
        r2.italic = True
        r3 = p.add_run("Hegde")
        r3.underline = True
        p.add_run(" holds 47.00% equity.")

    cfg = RedactionConfig(enable_ocr=False, salt="test-salt")
    res = Redactor(cfg).redact(_doc_bytes(build))
    doc = Document(io.BytesIO(res.output))
    runs = doc.paragraphs[0].runs

    assert len(runs) == 5
    text = "".join(r.text for r in runs)
    assert "Kushal" not in text
    assert "Hegde" not in text
    assert text.startswith("Promoter: ")
    assert text.endswith(" holds 47.00% equity.")
    # Check that individual styles survived
    assert runs[1].bold is True
    assert runs[2].italic is True
    assert runs[3].underline is True


def test_runs_formatting_not_wiped():
    """Verify that paragraph.text is NOT naively assigned, which would destroy run elements."""
    def build(d):
        p = d.add_paragraph()
        r = p.add_run("Confidential contact: ")
        r.font.name = "Arial"
        r2 = p.add_run("cs.connect@kshinternational.com")
        r2.bold = True

    cfg = RedactionConfig(enable_ocr=False)
    res = Redactor(cfg).redact(_doc_bytes(build))
    doc = Document(io.BytesIO(res.output))
    p = doc.paragraphs[0]
    assert len(p.runs) == 2
    assert p.runs[1].bold is True
    assert "kshinternational.com" not in p.runs[1].text
    assert "@" in p.runs[1].text


# ---------------------------------------------------------------------------
# 2. Deterministic Caching and Salt Registry Tests
# ---------------------------------------------------------------------------
def test_deterministic_pseudonym_cache():
    """Same entity yields the exact same pseudonym across multiple calls with same salt."""
    p1 = Pseudonymizer("seed-salt-123")
    p2 = Pseudonymizer("seed-salt-123")

    name = "Kushal Subbayya Hegde"
    p1.register(name, T.PERSON)
    p2.register(name, T.PERSON)

    fake1 = p1.fake_for(name, T.PERSON)
    fake2 = p2.fake_for(name, T.PERSON)
    assert fake1 == fake2
    assert fake1 != name

    # Test Anonymizer alias works identically
    anon = Anonymizer("seed-salt-123")
    anon.register(name, T.PERSON)
    assert anon.fake_for(name, T.PERSON) == fake1


def test_salt_isolation():
    """Different salts produce un-linkable, distinct pseudonyms."""
    p1 = Pseudonymizer("salt-alpha")
    p2 = Pseudonymizer("salt-beta")
    name = "Rajesh Kushal Hegde"

    p1.register(name, T.PERSON)
    p2.register(name, T.PERSON)

    assert p1.fake_for(name, T.PERSON) != p2.fake_for(name, T.PERSON)


def test_family_surname_consistency():
    """Family members sharing 'Hegde' receive the exact same synthetic surname."""
    p = Pseudonymizer("family-salt")
    p.register("Kushal Subbayya Hegde", T.PERSON)
    p.register("Pushpa Kushal Hegde", T.PERSON)
    p.register("Rajesh Kushal Hegde", T.PERSON)
    p.register("Rohit Kushal Hegde", T.PERSON)

    fake_kushal = p.fake_for("Kushal Subbayya Hegde", T.PERSON)
    fake_pushpa = p.fake_for("Pushpa Kushal Hegde", T.PERSON)
    fake_rajesh = p.fake_for("Rajesh Kushal Hegde", T.PERSON)
    fake_rohit = p.fake_for("Rohit Kushal Hegde", T.PERSON)

    sur_kushal = fake_kushal.split()[-1]
    sur_pushpa = fake_pushpa.split()[-1]
    sur_rajesh = fake_rajesh.split()[-1]
    sur_rohit = fake_rohit.split()[-1]

    # All family members must have identical synthetic surname
    assert sur_kushal == sur_pushpa == sur_rajesh == sur_rohit


# ---------------------------------------------------------------------------
# 3. Compiled Regex and Identifier Validation Tests
# ---------------------------------------------------------------------------
def test_cin_regex_and_pseudonym():
    """Validates Corporate Identity Number matching and format preservation."""
    cin = "U28129PN1979PLC141032"
    assert CIN_RE.search(cin) is not None

    p = Pseudonymizer("test")
    p.register(cin, T.CIN)
    fake = p.fake_for(cin, T.CIN)
    assert len(fake) == 21
    assert fake[0] == "U"
    assert fake[12:15] == "PLC"
    assert fake != cin


def test_pan_regex_and_holder_type():
    """Validates Permanent Account Number regex and 4th character (holder type) preservation."""
    pan = "NBWPS1951N"
    assert PAN_RE.search(pan) is not None

    p = Pseudonymizer("test")
    p.register(pan, T.PAN)
    fake = p.fake_for(pan, T.PAN)
    assert len(fake) == 10
    assert fake[3] == "P"  # 'P' for Individual/Person must be preserved
    assert fake != pan


def test_din_regex_and_pseudonym():
    """Validates Director Identification Number matching."""
    text = "DIN: 00135070"
    m = DIN_RE.search(text)
    assert m is not None
    assert m.group(1) == "00135070"

    p = Pseudonymizer("test")
    fake = p.fake_for("00135070", T.DIN)
    assert len(fake) == 8
    assert fake.startswith("00")


def test_sebi_reg_regex():
    """Validates SEBI Registration Numbers (INM, INR, INBI)."""
    assert SEBI_REG_RE.search("INM000013004") is not None
    assert SEBI_REG_RE.search("INR000004058") is not None
    assert SEBI_REG_RE.search("INBI00000063") is not None

    p = Pseudonymizer("test")
    fake_inbi = p.fake_for("INBI00000063", T.SEBI_REG)
    assert fake_inbi.startswith("INBI")
    assert len(fake_inbi) == 12


def test_aadhaar_verhoeff():
    """Validates 12-digit Aadhaar pattern and Verhoeff checksum algorithm."""
    aad = "2943 6593 3461"
    assert AADHAAR_RE.search(aad) is not None
    assert verhoeff_ok("294365933461") is True

    p = Pseudonymizer("test")
    p.register(aad, T.AADHAAR)
    fake = p.fake_for(aad, T.AADHAAR)
    assert len(fake) == 14
    assert verhoeff_ok("".join(c for c in fake if c.isdigit())) is True


def test_contact_details_and_email_domain():
    """Validates email and phone replacements."""
    p = Pseudonymizer("test")
    p.register("cs.connect@kshinternational.com", T.EMAIL)
    fake_email = p.fake_for("cs.connect@kshinternational.com", T.EMAIL)
    assert fake_email.startswith("cs.connect@")
    assert "kshinternational.com" not in fake_email
    assert fake_email.endswith(".example.com") or ".example" in fake_email

    p.register("+91 20 45053237", T.PHONE)
    fake_phone = p.fake_for("+91 20 45053237", T.PHONE)
    assert fake_phone.startswith("+91 ")


def test_statutory_text_allowlist():
    """Ensures stock exchange and regulatory bodies are not falsely redacted."""
    d = PIIDetector()
    text = "The Equity Shares will be listed on BSE Limited and National Stock Exchange of India Limited."
    hits = d.detect(text)
    assert hits == []


# ---------------------------------------------------------------------------
# 4. Table Context Traversal Tests
# ---------------------------------------------------------------------------
def test_table_cell_redaction():
    """Validates that table cells and nested structures are correctly processed."""
    def build(d):
        table = d.add_table(rows=2, cols=3)
        table.cell(0, 0).text = "Name"
        table.cell(0, 1).text = "DIN"
        table.cell(0, 2).text = "Address"
        table.cell(1, 0).text = "Sarthak Malvadkar"
        table.cell(1, 1).text = "00135070"
        table.cell(1, 2).text = "11/3 Village Birdewadi, Chakan, Pune - 410 501, Maharashtra, India"

    cfg = RedactionConfig(enable_ocr=False)
    res = Redactor(cfg).redact(_doc_bytes(build))
    doc = Document(io.BytesIO(res.output))
    row = doc.tables[0].rows[1]

    assert row.cells[0].text != "Sarthak Malvadkar"
    assert row.cells[1].text != "00135070"
    assert len(row.cells[1].text) == 8
    assert "410 501" not in row.cells[2].text


# ---------------------------------------------------------------------------
# 5. Image OCR and Media Masking Tests
# ---------------------------------------------------------------------------
@pytest.mark.skipif(not ImageRedactor.available(), reason="tesseract not installed")
def test_image_ocr_mask_and_banner_policy():
    """Tests image detection and redaction using both 'mask' (black rectangle) and 'banner'."""
    img = Image.new("RGB", (800, 400), "white")
    d = ImageDraw.Draw(img)
    try:
        f = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 24)
    except Exception:
        f = ImageFont.load_default()

    d.text((30, 20), "INCOME TAX DEPARTMENT", font=f, fill="black")
    d.text((30, 70), "Permanent Account Number Card", font=f, fill="black")
    d.text((30, 120), "NBWPS1951N", font=f, fill="black")
    d.text((30, 170), "Name: VISHAL SINGH", font=f, fill="black")
    d.text((30, 220), "Date of Birth: 06/05/2000", font=f, fill="black")

    b = io.BytesIO()
    img.save(b, "PNG")

    class _Part:
        partname = "/word/media/image_pan.png"
        blob = b.getvalue()

    # Test policy='mask' (solid black bounding boxes)
    p_mask = _Part()
    ps = Pseudonymizer()
    redactor_mask = ImageRedactor(PIIDetector(), ps, policy="mask")
    rep_mask = redactor_mask.process_part(p_mask)
    assert rep_mask.id_document is True
    assert len(rep_mask.hits) >= 2

    # Test policy='banner' ([CONFIDENTIAL ID CARD REDACTED])
    p_banner = _Part()
    redactor_banner = ImageRedactor(PIIDetector(), ps, policy="banner")
    rep_banner = redactor_banner.process_part(p_banner)
    assert rep_banner.id_document is True
