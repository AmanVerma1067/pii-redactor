import io

from docx import Document

from pii_redactor import RedactionConfig, Redactor
from pii_redactor.detector import PIIDetector
from pii_redactor.pseudonymizer import Pseudonymizer
from pii_redactor.validators import luhn_ok, verhoeff_ok
from pii_redactor.entities import EntityType as T


def _doc_bytes(build) -> bytes:
    d = Document()
    build(d)
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def test_split_runs_keep_formatting():
    def build(d):
        p = d.add_paragraph()
        p.add_run("Mr. ")
        r = p.add_run("Kus"); r.bold = True
        r = p.add_run("hal Subbayya "); r.italic = True
        p.add_run("Hegde")
        p.add_run(" is a Promoter.")
    out = Redactor(RedactionConfig(enable_ocr=False)).redact(_doc_bytes(build))
    runs = Document(io.BytesIO(out.output)).paragraphs[0].runs
    assert len(runs) == 5
    text = "".join(r.text for r in runs)
    assert "Kushal" not in text and "Hegde" not in text and text.startswith("Mr. ")
    assert runs[1].bold and runs[2].italic and runs[4].text == " is a Promoter."


def test_consistency_across_body_table_header():
    def build(d):
        d.sections[0].header.paragraphs[0].text = "KSH International Limited"
        d.add_paragraph("KSH International Limited is promoted by Mr. Rajesh Kushal Hegde.")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "Name", "DIN"
        t.cell(1, 0).text, t.cell(1, 1).text = "Rajesh Kushal Hegde", "00134926"
    res = Redactor(RedactionConfig(enable_ocr=False)).redact(_doc_bytes(build))
    doc = Document(io.BytesIO(res.output))
    body = doc.paragraphs[0].text
    header = doc.sections[0].header.paragraphs[0].text
    cell = doc.tables[0].cell(1, 0).text
    fake_org = header.replace(" International Limited", "")
    assert fake_org in body and "KSH" not in body
    assert cell in body                      # same person -> same pseudonym
    assert doc.tables[0].cell(1, 1).text != "00134926" and len(doc.tables[0].cell(1, 1).text) == 8


def test_detects_indian_identifiers():
    d = PIIDetector()
    text = ("CIN U28129PN1979PLC141032, PAN NBWPS1951N, Aadhaar 2943 6593 3461, SEBI INM000013004, "
            "DIN: 00135070, email cs.connect@kshinternational.com, Tel +91 22 4009 4400")
    types = {s.type for s in d.detect(text)}
    assert {T.CIN, T.PAN, T.AADHAAR, T.SEBI_REG, T.DIN, T.EMAIL, T.PHONE} <= types


def test_statutory_text_not_redacted():
    d = PIIDetector()
    text = ("Listed on BSE Limited and National Stock Exchange of India Limited under the SEBI ICDR Regulations. "
            "ISIN INE0ABC01018. 1,84,89,583 Equity Shares aggregating ₹7,100.00 million. Version 2.1.4.0.")
    assert d.detect(text) == []


def test_format_preserving_and_checksums():
    ps = Pseudonymizer("salt")
    aad = ps.fake_for("2943 6593 3461", T.AADHAAR)
    assert len(aad) == 14 and aad[4] == " " and verhoeff_ok(aad)
    cc = ps.fake_for("4111 1111 1111 1111", T.CREDIT_CARD)
    assert luhn_ok(cc) and cc.count(" ") == 3
    pan = ps.fake_for("NBWPS1951N", T.PAN)
    assert pan[3] == "P" and pan != "NBWPS1951N"
    assert ps.fake_for("+91 81081 14949", T.PHONE).startswith("+91 ")


def test_deterministic_with_salt():
    a, b, c = Pseudonymizer("x"), Pseudonymizer("x"), Pseudonymizer("y")
    for p in (a, b, c):
        p.register("Rashi Patil", T.PERSON)
    assert a.fake_for("Rashi Patil", T.PERSON) == b.fake_for("Rashi Patil", T.PERSON)
    assert a.fake_for("Rashi Patil", T.PERSON) != "Rashi Patil"


def test_email_follows_person_pseudonym():
    p = Pseudonymizer()
    p.register("Rashi Patil", T.PERSON)
    name = p.fake_for("Rashi Patil", T.PERSON).lower().split()
    email = p.fake_for("rashhi.patil@gmail.com", T.EMAIL)
    assert email == f"{name[0]}.{name[1]}@example.com"


def test_dob_rejects_impossible_years_and_phone_numbers():
    det = PIIDetector()
    dob = lambda text, **kw: [s.text for s in det.detect(text, **kw) if s.type == T.DOB]  # noqa: E731
    assert dob("Date of Birth: 06/05/2000") == ["06/05/2000"]
    assert dob("Tel: 91-20-2721 8080, Fax: 91-20-2721 8081", id_context=True) == []
    assert dob("DOB 12/12/2721") == []
    assert dob("DOB 12/12/1888") == []
    assert dob("born on 31/13/1990") == []
