"""Render EVALUATION.md as a Word document (headings, tables, lists, quotes, code) with python-docx only.

    python eval/md2docx.py EVALUATION.md Evaluation_Strategy_and_Metrics.docx
"""
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*)")
NUMERIC = re.compile(r"^[\d.,/%\s]+$")
BLOCK_START = re.compile(r"^(#|>|\||```|\s*([*-]|\d+\.)\s)")


def shade(cell, hex_fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(shd)


def add_inline(par, text, bold=False, size=None):
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            add_inline(par, tok[2:-2], bold=True, size=size)
            continue
        if tok.startswith("`") and tok.endswith("`"):
            r = par.add_run(tok[1:-1]); r.font.name = "Consolas"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            r.font.color.rgb = RGBColor(0xA3, 0x15, 0x15)
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = par.add_run(tok[1:-1]); r.italic = True
        else:
            r = par.add_run(tok)
        r.bold = bold or None
        if size:
            r.font.size = size


def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def add_table(doc, rows):
    header, body = rows[0], rows[2:]  # rows[1] is the |---| separator
    t = doc.add_table(rows=1 + len(body), cols=len(header))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate([header] + body):
        row = (row + [""] * len(header))[: len(header)]
        for ci, txt in enumerate(row):
            cell = t.cell(ri, ci); p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, txt, bold=(ri == 0), size=Pt(9))
            if ri == 0:
                shade(cell, "D9E2F3")
            elif txt.startswith("**"):
                shade(cell, "F2F2F2")
            if ri > 0 and NUMERIC.match(txt.replace("*", "")):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true")  # repeat header row across pages
    t.rows[0]._tr.get_or_add_trPr().append(h)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_code(doc, lines):
    t = doc.add_table(rows=1, cols=1); t.style = "Table Grid"
    cell = t.cell(0, 0); shade(cell, "F4F4F4")
    p = cell.paragraphs[0]
    for i, ln in enumerate(lines):
        r = p.add_run(ln); r.font.name = "Consolas"; r.font.size = Pt(8)
        if i < len(lines) - 1:
            r.add_break()
    doc.add_paragraph()


def convert(src: str, dst: str) -> None:
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2)
    normal = doc.styles["Normal"]; normal.font.name = "Calibri"; normal.font.size = Pt(10.5)

    lines = open(src, encoding="utf-8").read().splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("```"):
            j, block = i + 1, []
            while j < len(lines) and not lines[j].startswith("```"):
                block.append(lines[j]); j += 1
            add_code(doc, block); i = j + 1; continue
        if ln.lstrip().startswith("|"):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                rows.append(split_row(lines[i])); i += 1
            add_table(doc, rows); continue
        m = re.match(r"^(#{1,6})\s+(.*)", ln)
        if m:
            if len(m.group(1)) == 1:
                doc.add_heading(m.group(2), level=0)
            else:
                add_inline(doc.add_heading(level=len(m.group(1)) - 1), m.group(2))
            i += 1; continue
        if ln.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].startswith(">"):
                buf.append(lines[i].lstrip("> ").strip()); i += 1
            add_inline(doc.add_paragraph(style="Intense Quote"), " ".join(buf)); continue
        m = re.match(r"^(\s*)([*-]|\d+\.)\s+(.*)", ln)
        if m:
            buf = [m.group(3)]; i += 1
            while i < len(lines) and lines[i].startswith("  ") and not re.match(r"^\s*([*-]|\d+\.)\s", lines[i]):
                buf.append(lines[i].strip()); i += 1
            style = "List Number" if m.group(2)[0].isdigit() else "List Bullet"
            add_inline(doc.add_paragraph(style=style), " ".join(buf)); continue
        if not ln.strip():
            i += 1; continue
        buf = [ln.strip()]; i += 1
        while i < len(lines) and lines[i].strip() and not BLOCK_START.match(lines[i]):
            buf.append(lines[i].strip()); i += 1
        add_inline(doc.add_paragraph(), " ".join(buf))

    doc.core_properties.title = "Evaluation Strategy & Metric Report"
    doc.save(dst)


if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2])
    print("wrote", sys.argv[2])
